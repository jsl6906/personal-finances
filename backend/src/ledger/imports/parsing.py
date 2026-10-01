"""Reading uploaded spreadsheets into rows of strings, and parsing cell values."""

import csv
import hashlib
import io
import os
import re
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import PurePath

import duckdb
from dateutil import parser as dateparser

SPREADSHEET_EXT = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".xls"}
DOCUMENT_EXT = {".pdf", ".png", ".jpg", ".jpeg", ".webp", ".heic", ".gif", ".tif", ".tiff"}
DOCUMENT_MIME = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".heic": "image/heic",
    ".gif": "image/gif",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
}

FIELDS = (
    "txn_date",
    "posted_date",
    "description",
    "original_description",
    "amount",
    "debit",
    "credit",
    "category",
    "account",
    "notes",
    "check_number",
    "external_id",
    "ignore",
)


def detect_kind(filename: str) -> str | None:
    ext = PurePath(filename).suffix.lower()
    if ext in SPREADSHEET_EXT:
        return "spreadsheet"
    if ext in DOCUMENT_EXT:
        return "document"
    return None


def mime_for(filename: str, fallback: str | None) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext in DOCUMENT_MIME:
        return DOCUMENT_MIME[ext]
    if ext in (".csv", ".tsv", ".txt"):
        return "text/csv"
    if ext in (".xlsx", ".xlsm"):
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if ext == ".xls":
        return "application/vnd.ms-excel"
    return fallback or "application/octet-stream"


@dataclass
class Table:
    headers: list[str]
    rows: list[dict[str, str]]
    sheets: list[str] = field(default_factory=list)
    sheet: str | None = None


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _unique_headers(raw: list[str]) -> list[str]:
    out: list[str] = []
    for i, h in enumerate(raw):
        name = h.strip() or f"Column {i + 1}"
        base, n = name, 2
        while name in out:
            name = f"{base} ({n})"
            n += 1
        out.append(name)
    return out


def _from_matrix(matrix: list[list[str]]) -> tuple[list[str], list[dict[str, str]]]:
    """Find the header row (bank exports often have preamble lines) and build dict rows."""
    matrix = [r for r in matrix if any(c for c in r)]
    if not matrix:
        return [], []
    width = max(len(r) for r in matrix)
    header_idx = 0
    for i, r in enumerate(matrix[:25]):
        texts = [c for c in r if c and not _looks_numeric(c) and parse_date(c) is None]
        if len(texts) >= 2 and len([c for c in r if c]) >= max(2, int(width * 0.5)):
            header_idx = i
            break
    headers = _unique_headers([*(matrix[header_idx]), *[""] * (width - len(matrix[header_idx]))])
    rows = []
    for r in matrix[header_idx + 1 :]:
        r = [*r, *[""] * (width - len(r))]
        rows.append(dict(zip(headers, r, strict=False)))
    return headers, rows


def _read_csv(data: bytes) -> tuple[list[str], list[dict[str, str]]]:
    fd, path = tempfile.mkstemp(suffix=".csv")
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        try:
            con = duckdb.connect()
            rel = con.execute(
                "SELECT * FROM read_csv($p, all_varchar = true, header = false, auto_detect = true, "
                "null_padding = true, ignore_errors = true, sample_size = -1)",
                {"p": path},
            )
            matrix = [[_cell(c) for c in row] for row in rel.fetchall()]
            con.close()
            if matrix and max(len(r) for r in matrix) > 1:
                return _from_matrix(matrix)
        except duckdb.Error:
            pass
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|") if text.strip() else csv.excel
        return _from_matrix([[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)])
    finally:
        os.unlink(path)


def _read_xlsx(data: bytes, sheet: str | None) -> Table:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        names = wb.sheetnames
        ws = wb[sheet] if sheet in names else wb[names[0]]
        matrix = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
        headers, rows = _from_matrix(matrix)
        return Table(headers, rows, names, ws.title)
    finally:
        wb.close()


def _read_xls(data: bytes, sheet: str | None) -> Table:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    names = book.sheet_names()
    sh = book.sheet_by_name(sheet) if sheet in names else book.sheet_by_index(0)
    matrix = []
    for r in range(sh.nrows):
        row = []
        for c in range(sh.ncols):
            cell = sh.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(_cell(xlrd.xldate_as_datetime(cell.value, book.datemode)))
            else:
                row.append(_cell(cell.value))
        matrix.append(row)
    headers, rows = _from_matrix(matrix)
    return Table(headers, rows, names, sh.name)


def read_table(data: bytes, filename: str, sheet: str | None = None) -> Table:
    ext = PurePath(filename).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(data, sheet)
    if ext == ".xls":
        return _read_xls(data, sheet)
    headers, rows = _read_csv(data)
    return Table(headers, rows)


def header_signature(headers: list[str]) -> str:
    norm = "|".join(re.sub(r"\s+", " ", h.strip().lower()) for h in headers)
    return hashlib.sha256(norm.encode()).hexdigest()


def column_profile(headers: list[str], rows: list[dict[str, str]]) -> list[dict]:
    out = []
    for h in headers:
        samples = []
        for r in rows:
            v = r.get(h, "")
            if v and v not in samples:
                samples.append(v)
            if len(samples) == 3:
                break
        out.append({"name": h, "samples": samples})
    return out


_HEURISTICS: list[tuple[str, str]] = [
    (r"^(transaction |trans\.? |txn )?id$|transaction id|reference|^ref", "external_id"),
    (r"full description|original description", "original_description"),
    (r"post(ed|ing)? date", "posted_date"),
    (r"^(transaction |trans\.? )?date$|^date|trans(action)? date", "txn_date"),
    (r"desc|payee|merchant|memo|narrative|details|name", "description"),
    (r"debit|withdrawal|money out|charges?$|payments? out", "debit"),
    (r"credit|deposit|money in|payments? in", "credit"),
    (r"amount|amt|value", "amount"),
    (r"categor", "category"),
    (r"^account$|account name|^acct", "account"),
    (r"check|cheque", "check_number"),
    (r"note|comment", "notes"),
]


def heuristic_mapping(headers: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for h in headers:
        key = h.strip().lower()
        target = "ignore"
        for pattern, fld in _HEURISTICS:
            if re.search(pattern, key) and (fld not in used or fld in ("ignore",)):
                target = fld
                break
        mapping[h] = target
        if target != "ignore":
            used.add(target)
    return mapping


def mapping_complete(mapping: dict[str, str]) -> bool:
    fields = set(mapping.values())
    has_amount = "amount" in fields or ("debit" in fields or "credit" in fields)
    return "txn_date" in fields and has_amount and ("description" in fields or "original_description" in fields)


def _looks_numeric(s: str) -> bool:
    return parse_amount(s) is not None


_AMOUNT_CLEAN = re.compile(r"[$€£,\s]|USD", re.IGNORECASE)


def parse_amount(value: str | None) -> Decimal | None:
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1]
    upper = s.upper()
    if upper.endswith("CR"):
        s = s[:-2]
    elif upper.endswith("DR"):
        negative, s = True, s[:-2]
    s = _AMOUNT_CLEAN.sub("", s).replace("−", "-")
    if s.endswith("-"):
        negative, s = True, s[:-1]
    if s.startswith("+"):
        s = s[1:]
    if not re.fullmatch(r"-?\d+(\.\d+)?", s):
        return None
    try:
        d = Decimal(s)
    except InvalidOperation:
        return None
    return (-d if negative else d).quantize(Decimal("0.01"))


_EXCEL_EPOCH = date(1899, 12, 30)
_DATE_HINT = re.compile(
    r"\d{1,4}[-/.]\d{1,2}([-/.]\d{1,4})?|[A-Za-z]{3,9}\.? \d{1,2},? \d{2,4}|\d{1,2} [A-Za-z]{3,9}\.? \d{2,4}"
)


def parse_date(value: str | None, fmt: str | None = None, dayfirst: bool = False) -> date | None:
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    if fmt and fmt != "auto":
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            return None
    if re.fullmatch(r"\d{5}(\.\d+)?", s):  # Excel serial day number
        n = float(s)
        return _EXCEL_EPOCH + timedelta(days=int(n)) if 20000 < n < 80000 else None
    if not _DATE_HINT.search(s):
        return None
    try:
        d = dateparser.parse(s, dayfirst=dayfirst, default=datetime(2000, 1, 1)).date()
    except (ValueError, OverflowError):
        return None
    return d if 1970 <= d.year <= 2100 else None
